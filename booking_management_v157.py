"""Operational booking coverage report, week-over-week comparisons and safe exports."""
from datetime import date
from pathlib import Path
from tempfile import NamedTemporaryFile
import os
from booking_calendar_v156 import calendar_weeks, coverage_band, coverage_insights


def weekly_coverage(rows, *, week_start=5):
    """Average booked room coverage per calendar week; partial weeks use observed days."""
    result=[]
    for week in calendar_weeks(rows,week_start=week_start):
        observed=[x for x in week if x is not None]
        avg=round(sum(x['booked_pct'] for x in observed)/len(observed),1)
        result.append({'start':observed[0]['date'],'end':observed[-1]['date'],
                       'days':len(observed),'average_booked_pct':avg,
                       'change_pp':round(avg-result[-1]['average_booked_pct'],1) if result else None})
    return result


def report_rows(rows):
    """Validate and return a safe bounded report with no personally identifiable data."""
    if not isinstance(rows,list) or len(rows)>90:raise ValueError('Invalid booking report')
    calendar_weeks(rows)
    result=[]
    for row in rows:
        if type(row.get('booked_rooms')) is not int or type(row.get('total_rooms')) is not int or not 0<=row['booked_rooms']<=row['total_rooms']:
            raise ValueError('Invalid room counts')
        band=coverage_band(row)
        result.append((date.fromisoformat(row['date']).isoformat(),row['booked_rooms'],row['total_rooms'],row['booked_pct'],band))
    return result


def _atomic(target, suffix, write):
    target=Path(target)
    if target.suffix.lower()!=suffix:raise ValueError('Incorrect export extension')
    temp=None
    try:
        with NamedTemporaryFile(dir=target.parent,suffix=suffix,delete=False) as f:temp=Path(f.name)
        write(temp)
        os.replace(temp,target)
    finally:
        if temp is not None:temp.unlink(missing_ok=True)
    return target


def export_booking_excel(rows,path,*,fa=False):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill,Font,Alignment
    values=report_rows(rows);summary=coverage_insights(rows)
    def write(temp):
        wb=Workbook();ws=wb.active;ws.title='Booking coverage'
        ws.append(('Date','Booked rooms','Total rooms','Booked coverage %','Band'))
        for row in values:ws.append(row)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='17364B');cell.font=Font(color='FFFFFF',bold=True)
        bands={'low':'FFF1D6','medium':'D9F3F1','high':'D8F2E2','unavailable':'E5E7EB'}
        for i,row in enumerate(values,2):
            for cell in ws[i]:cell.fill=PatternFill('solid',fgColor=bands[row[4]])
            ws.cell(i,4).number_format='0.0'
        for col,width in {'A':18,'B':18,'C':18,'D':23,'E':18}.items():ws.column_dimensions[col].width=width
        weeks=wb.create_sheet('Weekly comparison');weeks.append(('Start','End','Days','Average booked %','Change (pp)'))
        for item in weekly_coverage(rows,week_start=5 if fa else 0):
            weeks.append((item['start'],item['end'],item['days'],item['average_booked_pct'],item['change_pp']))
        stats=wb.create_sheet('Summary');stats.append(('Metric','Value'))
        for key,value in summary.items():stats.append((key,value))
        wb.save(temp)
    return _atomic(path,'.xlsx',write)


def export_booking_pdf(rows,path,*,fa=False):
    from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A4,landscape
    values=report_rows(rows);summary=coverage_insights(rows)
    def write(temp):
        doc=SimpleDocTemplate(str(temp),pagesize=landscape(A4),rightMargin=28,leftMargin=28)
        styles=getSampleStyleSheet();story=[Paragraph('Booking coverage — recorded reservations',styles['Title']),Spacer(1,10)]
        story.append(Paragraph(f"Days: {summary['days']} | Low: {summary['low']} | Medium: {summary['medium']} | High: {summary['high']}",styles['Normal']))
        story.append(Spacer(1,12))
        data=[['Date','Booked rooms','Total rooms','Coverage %','Band']]+[[str(v) for v in row] for row in values]
        table=Table(data,repeatRows=1,colWidths=[145,110,110,120,100]);table.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#17364B')),
            ('TEXTCOLOR',(0,0),(-1,0),colors.white),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#EDF5F8')]),
            ('BOTTOMPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7)]))
        story.append(table);doc.build(story)
    return _atomic(path,'.pdf',write)
