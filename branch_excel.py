"""Native Excel branch reservation export with safe text cells."""
from pathlib import Path
from branch_csv import HEADERS, safe_cell


def export_branch_workbook(path, rows, language='fa', branch_name='', date_from='', date_to=''):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    if language not in HEADERS:
        raise ValueError('Unsupported language')
    path = Path(path)
    records = list(rows)
    if any(len(record) != 5 for record in records):
        raise ValueError('Expected five reservation fields')
    wb = Workbook()
    ws = wb.active
    ws.title = 'رزروها' if language == 'fa' else 'Reservations'
    ws.sheet_view.rightToLeft = language == 'fa'
    ws.merge_cells('A1:E1')
    title = ws['A1']
    title.value = 'گزارش رزروهای شعبه' if language == 'fa' else 'Branch Reservations Report'
    title.font = Font(name='Calibri', size=17, bold=True, color='FFFFFF')
    title.fill = PatternFill('solid', fgColor='142D4E')
    title.alignment = Alignment(horizontal='right' if language == 'fa' else 'left', vertical='center')
    ws.row_dimensions[1].height = 34
    ws.merge_cells('A2:E2')
    ws['A2'] = f"{branch_name}  |  {date_from} — {date_to}"
    ws['A2'].alignment = Alignment(horizontal='right' if language == 'fa' else 'left')
    ws.row_dimensions[2].height = 25
    for column, header in enumerate(HEADERS[language], 1):
        cell = ws.cell(4, column, header)
        cell.fill = PatternFill('solid', fgColor='087E8B')
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[4].height = 27
    for index, record in enumerate(records, 5):
        for col, value in enumerate(record, 1):
            cell = ws.cell(index, col)
            # All imported fields are literal text: never Excel formulas.
            cell.value = safe_cell(value)
            cell.data_type = 's'
            cell.alignment = Alignment(horizontal='right' if language == 'fa' else 'left', vertical='center')
            if index % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='EAF4F5')
        ws.row_dimensions[index].height = 22
    for index, width in enumerate((19, 19, 32, 25, 25), 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f'A4:E{max(4, 4 + len(records))}'
    wb.save(path)
    return path
