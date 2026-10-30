"""Authorized multi-branch comparison with explicit non-ADR financial metrics."""
from datetime import date
from pathlib import Path
from branch_reports import branch_summary
from branch_csv import safe_cell


def branch_kpis(summary, start, end):
    days = (date.fromisoformat(end) - date.fromisoformat(start)).days
    if days <= 0:
        raise ValueError('Invalid reporting interval')
    rooms = int(summary['rooms'])
    nights = int(summary['occupied_nights'])
    payments = int(summary['payments'])
    if min(rooms, nights, payments) < 0 or nights > rooms * days:
        raise ValueError('Invalid occupancy metrics')
    available = rooms * days
    return dict(available_nights=available, occupancy_percent=round(100 * nights / available, 2) if available else 0.0,
                collected_per_occupied_night=round(payments / nights, 2) if nights else 0.0)


def authorized_comparison(conn, *, user_id, role, start, end):
    """Only return branches explicitly visible to the authenticated user."""
    from branch_access import allowed_branches, require_branch
    result = []
    for branch_id, name in allowed_branches(conn, user_id, role):
        require_branch(conn, user_id, role, branch_id)
        summary = branch_summary(conn, branch_id, start, end)
        result.append(dict(branch_id=branch_id, branch_name=name, **summary, **branch_kpis(summary, start, end)))
    return result


def export_comparison_workbook(path, data, *, start, end, language='fa'):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment
    if language not in ('fa', 'en'):
        raise ValueError('Unsupported language')
    if date.fromisoformat(start) >= date.fromisoformat(end):
        raise ValueError('Invalid reporting interval')
    fa = language == 'fa'
    wb = Workbook(); ws = wb.active
    ws.title = 'مقایسه شعب' if fa else 'Branch Comparison'
    ws.sheet_view.rightToLeft = fa
    ws.merge_cells('A1:H1'); ws['A1'] = 'مقایسه عملکرد شعب' if fa else 'Branch Performance Comparison'
    ws['A1'].font = Font(size=16, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor='142D4E')
    ws['A2'] = f'{start} — {end} (end exclusive)'
    columns = ('شعبه','اتاق فعال','رزرو','شب‌اتاق اشغال','شب‌اتاق قابل فروش','درصد اشغال','دریافتی ثبت‌شده','دریافتی به ازای شب‌اتاق اشغال') if fa else ('Branch','Active rooms','Reservations','Occupied room-nights','Available room-nights','Occupancy %','Recorded payments','Collected per occupied room-night')
    for col, name in enumerate(columns, 1):
        cell=ws.cell(4,col,name);cell.fill=PatternFill('solid',fgColor='087E8B');cell.font=Font(bold=True,color='FFFFFF')
    for row_idx, record in enumerate(data, 5):
        ws.cell(row_idx,1,safe_cell(record['branch_name'])).data_type='s'
        values=(record['rooms'],record['reservations'],record['occupied_nights'],record['available_nights'],record['occupancy_percent'],record['payments'],record['collected_per_occupied_night'])
        for col,value in enumerate(values,2):
            cell=ws.cell(row_idx,col,value);cell.number_format='0.00' if col in (6,8) else '#,##0'
            if row_idx%2==0:cell.fill=PatternFill('solid',fgColor='EAF4F5')
    ws.freeze_panes='B5';ws.auto_filter.ref=f'A4:H{max(4,4+len(data))}'
    for col in 'ABCDEFGH':ws.column_dimensions[col].width=26 if col!='H' else 43
    if data:
        chart=BarChart();chart.title='درصد اشغال' if fa else 'Occupancy (%)'
        chart.y_axis.title='%';chart.add_data(Reference(ws,min_col=6,min_row=4,max_row=4+len(data)),titles_from_data=True)
        chart.set_categories(Reference(ws,min_col=1,min_row=5,max_row=4+len(data)))
        ws.add_chart(chart,'J5')
    ws['A3']='دریافتی / شب‌اتاق اشغال ADR یا RevPAR نیست.' if fa else 'Collected per occupied room-night is NOT ADR or RevPAR.'
    output=Path(path);wb.save(output);return output
