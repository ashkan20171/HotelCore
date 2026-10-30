"""Admin-only system-health reporting. Read-only; exports never create health checks."""
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from user_access_management import _actor
from system_health_v119 import migrate


def report_data(conn, actor_id, *, limit=100):
    _actor(conn, actor_id)
    if type(limit) is not int or not 1 <= limit <= 500:
        raise ValueError('limit must be 1..500')
    migrate(conn)
    records = conn.execute('''SELECT checked_at,database_status,monitor_state,failure_count
        FROM system_health_checks ORDER BY id DESC LIMIT ?''', (limit,)).fetchall()
    total = len(records)
    warnings = sum(1 for r in records if r[1] != 'ok')
    latest = records[0] if records else None
    last_ok = next((r[0] for r in records if r[1] == 'ok'), None)
    grouped = defaultdict(lambda: [0, 0, 0])
    for checked_at, status, state, failures in records:
        day = checked_at[:10]
        grouped[day][0] += 1
        grouped[day][1] += int(status != 'ok')
        grouped[day][2] = max(grouped[day][2], int(failures))
    trends = [(day, *grouped[day]) for day in sorted(grouped)]
    if latest is None:
        indicator = 'no_data'
    elif latest[1] != 'ok' or latest[2] in ('error', 'permission_error', 'failed', 'stopped'):
        indicator = 'warning'
    else:
        indicator = 'healthy'
    return {'total': total, 'warnings': warnings, 'healthy': total - warnings,
            'warning_rate': round(100 * warnings / total, 1) if total else 0.0,
            'latest': latest, 'last_healthy_check': last_ok,
            'indicator': indicator, 'trends': trends, 'records': records}


def _safe_cell(value):
    """Prevent spreadsheet formula injection in data sourced from database strings."""
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def write_excel(path, conn, actor_id, *, limit=100, language='fa'):
    if language not in ('fa', 'en'):
        raise ValueError('Unsupported language')
    data = report_data(conn, actor_id, limit=limit)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, Reference
    wb = Workbook()
    summary = wb.active
    summary.title = 'Summary'
    labels = [('وضعیت سلامت' if language == 'fa' else 'Health indicator', data['indicator']),
              ('تعداد بررسی‌ها' if language == 'fa' else 'Checks', data['total']),
              ('هشدارها' if language == 'fa' else 'Warnings', data['warnings']),
              ('درصد هشدار' if language == 'fa' else 'Warning rate (%)', data['warning_rate']),
              ('آخرین بررسی سالم' if language == 'fa' else 'Last healthy check', data['last_healthy_check'] or '—')]
    for row in labels:
        summary.append([_safe_cell(v) for v in row])
    details = wb.create_sheet('History')
    details.append(['UTC timestamp', 'Database status', 'Monitor state', 'Failure count'])
    for row in data['records']:
        details.append([_safe_cell(v) for v in row])
    trends = wb.create_sheet('Daily trends')
    trends.append(['UTC date', 'Checks', 'Warnings', 'Max failure count'])
    for row in data['trends']:
        trends.append([_safe_cell(v) for v in row])
    if data['trends']:
        chart = BarChart()
        chart.title = 'Health checks and warnings'
        chart.add_data(Reference(trends, min_col=2, max_col=3, min_row=1, max_row=trends.max_row), titles_from_data=True)
        chart.set_categories(Reference(trends, min_col=1, min_row=2, max_row=trends.max_row))
        trends.add_chart(chart, 'F2')
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.sheet_view.rightToLeft = language == 'fa'
        sheet.column_dimensions['A'].width = 34
        sheet.column_dimensions['B'].width = 25
        sheet.column_dimensions['C'].width = 24
        sheet.column_dimensions['D'].width = 20
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='15344B')
            cell.font = Font(bold=True, color='FFFFFF')
    wb.save(str(path))
    return data['total']


def write_pdf(path, conn, actor_id, *, limit=100, language='fa'):
    if language not in ('fa', 'en'):
        raise ValueError('Unsupported language')
    data = report_data(conn, actor_id, limit=limit)
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT, TA_LEFT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font = 'Helvetica'
    for candidate in ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
                      'C:/Windows/Fonts/tahoma.ttf', 'C:/Windows/Fonts/arial.ttf'):
        if Path(candidate).exists():
            try:
                pdfmetrics.registerFont(TTFont('HotelHealthUnicode', candidate))
                font = 'HotelHealthUnicode'
                break
            except Exception:
                pass
    def localize(value):
        value = str(value)
        if language == 'fa':
            try:
                import arabic_reshaper
                from bidi.algorithm import get_display
                return get_display(arabic_reshaper.reshape(value))
            except ImportError:
                pass
        return value
    style = ParagraphStyle('health', fontName=font, fontSize=9, leading=14,
                           alignment=TA_RIGHT if language == 'fa' else TA_LEFT)
    title_style = ParagraphStyle('health_title', parent=style, fontSize=16, leading=24)
    title = 'گزارش سلامت سیستم' if language == 'fa' else 'System Health Report'
    story = [Paragraph(localize(title), title_style), Spacer(1, 12)]
    for key, value in [('Indicator', data['indicator']), ('Checks', data['total']),
                       ('Warnings', data['warnings']), ('Warning rate', f"{data['warning_rate']}%"),
                       ('Last healthy check', data['last_healthy_check'] or '—')]:
        story.append(Paragraph(localize(f'{key}: {value}'), style))
    story.append(Spacer(1, 16))
    rows = [['UTC date', 'Checks', 'Warnings', 'Max failures']]
    rows.extend([[str(v) for v in record] for record in data['trends']])
    table = Table(rows, repeatRows=1, colWidths=[180, 100, 100, 115])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#15344B')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, -1), font),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
        ('GRID', (0, 0), (-1, -1), .3, colors.HexColor('#C5D1D9')),
        ('TOPPADDING', (0, 0), (-1, -1), 7),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 7)]))
    story.append(table)
    story.append(Spacer(1, 12))
    story.append(Paragraph(localize('Historical observations only; not a guarantee of system health.'), style))
    SimpleDocTemplate(str(path), pagesize=(595, 842), leftMargin=40, rightMargin=40).build(story)
    return data['total']
