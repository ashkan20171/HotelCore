import unittest,tempfile
from pathlib import Path
from executive_dashboard_export_v133 import export_excel,export_pdf,validate

DATA={'room_types':[{'type':'Suite','rooms':2,'booked_nights':4,'booked_value':600,'estimated_adr':150,'estimated_revpar':100,'occupancy_pct':66.7}], 'series':[{'date':'2026-10-20','income':500,'expenses':200}]}
class ExportTests(unittest.TestCase):
    def test_invalid(self):
        with self.assertRaises(ValueError):validate({'room_types':[]})
    def test_excel(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'out.xlsx';export_excel(DATA,p);w=load_workbook(p);self.assertEqual(w.worksheets[0]['A2'].value,'Suite');self.assertEqual(w.worksheets[1]['B2'].value,500)
    def test_pdf(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'out.pdf';export_pdf(DATA,p);self.assertTrue(p.read_bytes().startswith(b'%PDF'))
    def test_formula_protection(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            data=dict(DATA,room_types=[dict(DATA['room_types'][0],type='=1+1')]);p=Path(d)/'out.xlsx';export_excel(data,p)
            self.assertEqual(load_workbook(p).active['A2'].data_type,'s')
