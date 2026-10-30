import sqlite3
import tempfile
import unittest
from pathlib import Path
from system_health_v119 import snapshot
from system_health_reports_v120 import report_data, write_excel, write_pdf, _safe_cell

class HealthReportTests(unittest.TestCase):
    def setUp(self):
        self.db = sqlite3.connect(':memory:')
        self.db.execute('CREATE TABLE users(id INTEGER PRIMARY KEY,role TEXT,active INTEGER)')
        self.db.executemany('INSERT INTO users VALUES(?,?,?)', [(1,'admin',1),(2,'reception',1)])
    def tearDown(self):self.db.close()
    def test_empty_report(self):
        result=report_data(self.db,1)
        self.assertEqual((result['total'],result['indicator'],result['warning_rate']), (0,'no_data',0))
    def test_trend_aggregation(self):
        snapshot(self.db,1,monitor_state='healthy')
        snapshot(self.db,1,monitor_state='error')
        self.db.execute("UPDATE system_health_checks SET database_status='warning' WHERE id=1")
        result=report_data(self.db,1)
        self.assertEqual((result['total'],result['warnings'],result['warning_rate']), (2,1,50.0))
        self.assertEqual(result['trends'][0][1:3],(2,1))
    def test_authorization(self):
        with self.assertRaises(PermissionError):report_data(self.db,2)
    def test_limit_validation(self):
        for limit in (True,0,501,'10'):
            with self.assertRaises(ValueError):report_data(self.db,1,limit=limit)
    def test_excel_export(self):
        from openpyxl import load_workbook
        snapshot(self.db,1)
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'health.xlsx'
            self.assertEqual(write_excel(path,self.db,1,language='en'),1)
            book=load_workbook(path)
            self.assertEqual(book['History'].max_row,2)
            self.assertEqual(book['Daily trends'].max_row,2)
    def test_pdf_export(self):
        snapshot(self.db,1)
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'health.pdf'
            self.assertEqual(write_pdf(path,self.db,1,language='en'),1)
            self.assertTrue(path.read_bytes().startswith(b'%PDF'))
    def test_formula_protection(self):
        self.assertEqual(_safe_cell('=HYPERLINK("x")'), "'=HYPERLINK(\"x\")")
        self.assertEqual(_safe_cell(' @SUM(A1)'), "' @SUM(A1)")
        self.assertEqual(_safe_cell('healthy'),'healthy')
    def test_invalid_language(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):write_excel(Path(tmp)/'x.xlsx',self.db,1,language='de')
            with self.assertRaises(ValueError):write_pdf(Path(tmp)/'x.pdf',self.db,1,language='de')
