import sqlite3,unittest,tempfile
from pathlib import Path
from datetime import date
from unittest.mock import patch
from executive_dashboard_v131 import snapshot
from executive_dashboard_export_v133 import export_excel

class Stage141Tests(unittest.TestCase):
    def setUp(self):
        self.db=sqlite3.connect(':memory:')
        self.db.executescript("""CREATE TABLE rooms(id INTEGER PRIMARY KEY);
        CREATE TABLE reservations(id INTEGER PRIMARY KEY,room_id INTEGER,checkin TEXT,checkout TEXT,status TEXT,created_at TEXT);
        CREATE TABLE payments(id INTEGER PRIMARY KEY,amount INTEGER,created_at TEXT);
        CREATE TABLE expenses(id INTEGER PRIMARY KEY,amount INTEGER,created_at TEXT);
        INSERT INTO rooms VALUES(1); INSERT INTO rooms VALUES(2);
        INSERT INTO reservations VALUES(1,1,'2026-10-22','2026-10-24','confirmed','2026-10-23 12:00:00');
        INSERT INTO reservations VALUES(2,1,'2026-10-22','2026-10-24','confirmed','2026-10-23 12:00:00');
        INSERT INTO reservations VALUES(3,2,'2026-10-22','2026-10-24','cancelled','2026-10-23 12:00:00');""")
    def tearDown(self):self.db.close()
    def test_no_double_counting(self):
        with patch('executive_dashboard_v131._actor'):
            data=snapshot(self.db,1,days=2,today=date(2026,10,23))
        self.assertEqual(data['current']['occupied_nights'],2)
        self.assertEqual(data['current']['occupancy_pct'],50)
    def test_no_room_counts(self):
        self.db.execute('DELETE FROM rooms')
        with patch('executive_dashboard_v131._actor'):
            data=snapshot(self.db,1,days=2,today=date(2026,10,23))
        self.assertEqual(data['current']['occupied_nights'],0)
    def test_excel_formatting(self):
        from openpyxl import load_workbook
        data={'room_types':[{'type':'Suite','rooms':2,'booked_nights':3,'booked_value':500,'estimated_adr':166,'estimated_revpar':100,'occupancy_pct':50}],
              'series':[{'date':'2026-10-23','income':100,'expenses':30}]}
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'report.xlsx';export_excel(data,path)
            wb=load_workbook(path);ws=wb.active
            self.assertEqual(ws.freeze_panes,'A2')
            self.assertTrue(ws.auto_filter.ref)
            self.assertEqual(ws['A1'].font.color.rgb,'00FFFFFF')
            self.assertEqual(wb['Daily cash flow']['B2'].value,100)
