import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from branch_comparison import branch_kpis, authorized_comparison, export_comparison_workbook

class KPITests(unittest.TestCase):
    def test_occupancy(self):
        self.assertEqual(branch_kpis({'rooms':10,'occupied_nights':15,'payments':3000},'2026-01-01','2026-01-04')['occupancy_percent'],50)
    def test_collected_not_adr(self):
        self.assertEqual(branch_kpis({'rooms':10,'occupied_nights':15,'payments':3000},'2026-01-01','2026-01-04')['collected_per_occupied_night'],200)
    def test_empty_rooms(self):
        self.assertEqual(branch_kpis({'rooms':0,'occupied_nights':0,'payments':0},'2026-01-01','2026-01-02')['occupancy_percent'],0)
    def test_overbooking_invalid(self):
        with self.assertRaises(ValueError):branch_kpis({'rooms':1,'occupied_nights':3,'payments':0},'2026-01-01','2026-01-02')
    def test_invalid_interval(self):
        with self.assertRaises(ValueError):branch_kpis({'rooms':1,'occupied_nights':0,'payments':0},'2026-01-02','2026-01-01')
    def test_negative_payments_invalid(self):
        with self.assertRaises(ValueError):branch_kpis({'rooms':1,'occupied_nights':0,'payments':-1},'2026-01-01','2026-01-02')
    def test_access_scope(self):
        summary={'rooms':2,'reservations':1,'occupied_nights':1,'payments':50}
        with patch('branch_access.allowed_branches',return_value=[(7,'Authorized')]), patch('branch_access.require_branch') as auth, patch('branch_comparison.branch_summary',return_value=summary) as report:
            result=authorized_comparison(object(),user_id=3,role='staff',start='2026-01-01',end='2026-01-02')
        self.assertEqual(len(result),1);self.assertEqual(result[0]['branch_id'],7);auth.assert_called_once();report.assert_called_once()

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook=None
@unittest.skipUnless(load_workbook,'openpyxl missing')
class ExcelTests(unittest.TestCase):
    def test_bilingual_and_safe(self):
        with tempfile.TemporaryDirectory() as temp:
            for language in ('fa','en'):
                path=Path(temp)/(language+'.xlsx')
                row=dict(branch_name='=evil',rooms=2,reservations=1,occupied_nights=1,available_nights=2,occupancy_percent=50,payments=200,collected_per_occupied_night=200)
                export_comparison_workbook(path,[row],start='2026-01-01',end='2026-01-02',language=language)
                sheet=load_workbook(path).active
                self.assertEqual(sheet['A5'].data_type,'s');self.assertTrue(sheet['A5'].value.startswith("'="))
                self.assertEqual(sheet.sheet_view.rightToLeft,language=='fa');self.assertEqual(len(sheet._charts),1)
    def test_bad_language(self):
        with self.assertRaises(ValueError):export_comparison_workbook('no.xlsx',[],start='2026-01-01',end='2026-01-02',language='xx')
