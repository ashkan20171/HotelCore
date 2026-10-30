import tempfile
import unittest
from pathlib import Path
from branch_excel import export_branch_workbook
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

@unittest.skipUnless(load_workbook, 'openpyxl is not installed')
class ExcelExportTests(unittest.TestCase):
    def make_file(self, rows, language='fa'):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        path = Path(self.tmp.name) / 'report.xlsx'
        export_branch_workbook(path, rows, language, 'Branch 1', '2026-01-01', '2026-02-01')
        return load_workbook(path).active

    def test_fa_headers_and_direction(self):
        ws = self.make_file([('1', '101', 'علی', '2026-01-01', '2026-01-02')])
        self.assertTrue(ws.sheet_view.rightToLeft)
        self.assertEqual(ws['C5'].value, 'علی')
        self.assertEqual(ws['A4'].value, 'شناسه رزرو')

    def test_en_headers(self):
        ws = self.make_file([], 'en')
        self.assertFalse(ws.sheet_view.rightToLeft)
        self.assertEqual(ws['A4'].value, 'Reservation ID')

    def test_formula_injection_is_text(self):
        ws = self.make_file([('1', '101', '=HYPERLINK("evil")', '2026-01-01', '2026-01-02')])
        self.assertEqual(ws['C5'].data_type, 's')
        self.assertTrue(ws['C5'].value.startswith("'="))

    def test_freeze_and_filter(self):
        ws = self.make_file([('1','101','Guest','2026-01-01','2026-01-02')])
        self.assertEqual(ws.freeze_panes, 'A5')
        self.assertEqual(ws.auto_filter.ref, 'A4:E5')

    def test_reject_invalid_language(self):
        with self.assertRaises(ValueError):
            export_branch_workbook('bad.xlsx', [], 'xx')

    def test_reject_malformed_rows(self):
        with self.assertRaises(ValueError):
            export_branch_workbook('bad.xlsx', [('1','2')])

if __name__ == '__main__':
    unittest.main()
