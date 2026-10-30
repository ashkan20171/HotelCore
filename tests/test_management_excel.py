import tempfile
import unittest
from pathlib import Path
from management_excel import payment_series, export_management_workbook
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

class SeriesTests(unittest.TestCase):
    def test_monthly_aggregation(self):
        data = [(1,1,100,'نقدی','2026-01-10'), (2,1,50,'card','2026-01-20'), (3,1,75,'cash','2026-02-03')]
        self.assertEqual(payment_series(data,'2026-01-01','2026-03-01'), [('2026-01',150,2),('2026-02',75,1)])
    def test_excludes_end(self):
        self.assertEqual(payment_series([(1,1,10,'cash','2026-02-01')],'2026-01-01','2026-02-01'), [])
    def test_bad_interval(self):
        with self.assertRaises(ValueError):payment_series([], '2026-01-01', '2026-01-01')
    def test_bad_payment(self):
        with self.assertRaises(ValueError):payment_series([(1,2)], '2026-01-01', '2026-02-01')

@unittest.skipUnless(load_workbook, 'openpyxl not installed')
class WorkbookTests(unittest.TestCase):
    def build(self,language='fa', payments=None):
        tmp=tempfile.TemporaryDirectory();self.addCleanup(tmp.cleanup)
        path=Path(tmp.name)/'report.xlsx'
        export_management_workbook(path,branch_name='تهران',start='2026-01-01',end='2026-02-01',summary={'rooms':10,'reservations':1,'occupied_nights':3,'payments':250},reservations=[(1,'101','=DANGER','2026-01-01','2026-01-04')],payments=payments if payments is not None else [(1,1,250,'card','2026-01-04')],language=language)
        return load_workbook(path)
    def test_four_sheets(self):
        self.assertEqual(len(self.build().sheetnames),4)
    def test_fa_rtl(self):
        self.assertTrue(all(ws.sheet_view.rightToLeft for ws in self.build()))
    def test_en_ltr(self):
        self.assertTrue(all(not ws.sheet_view.rightToLeft for ws in self.build('en')))
    def test_formula_safe(self):
        cell=self.build()['رزروها']['C5'];self.assertEqual(cell.data_type,'s');self.assertTrue(cell.value.startswith("'="))
    def test_numeric_amount(self):
        self.assertEqual(self.build()['پرداخت‌ها']['C5'].value,250)
    def test_chart_present(self):
        self.assertEqual(len(self.build()['روند درآمد']._charts),1)
    def test_no_chart_without_data(self):
        self.assertEqual(len(self.build(payments=[])['روند درآمد']._charts),0)
    def test_bad_language(self):
        with self.assertRaises(ValueError):export_management_workbook('unused.xlsx',branch_name='',start='2026-01-01',end='2026-02-01',summary={},reservations=[],payments=[],language='de')
