"""Read-only export of a previously authorized executive dashboard snapshot."""
from pathlib import Path
from tempfile import NamedTemporaryFile
import os

FIELDS=('type','rooms','booked_nights','booked_value','estimated_adr','estimated_revpar','occupancy_pct')
HEADERS=('Room type','Rooms','Booked nights','Estimated booked value','Estimated ADR','Estimated RevPAR','Occupancy %')

def _safe(value):
    text=str(value if value is not None else '')
    return "'"+text if text.lstrip().startswith(('=','+','-','@','\t','\r')) else text

def validate(data):
    if not isinstance(data,dict) or not isinstance(data.get('room_types'),list) or not isinstance(data.get('series'),list):
        raise ValueError('Invalid dashboard snapshot')
    if len(data['room_types'])>10000 or len(data['series'])>90:
        raise ValueError('Snapshot too large')
    for row in data['room_types']:
        if not isinstance(row,dict) or not all(k in row for k in FIELDS):
            raise ValueError('Invalid room type row')
    for row in data['series']:
        if not isinstance(row,dict) or not all(k in row for k in ('date','income','expenses')):
            raise ValueError('Invalid trend row')

def export_excel(data,path):
    """Atomic xlsx output; callers must obtain authorized data before calling."""
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    validate(data)
    target=Path(path)
    if target.suffix.lower()!='.xlsx':raise ValueError('Expected .xlsx output')
    wb=Workbook(); ws=wb.active;ws.title='Room performance';ws.append(HEADERS)
    for row in data['room_types']:
        ws.append([_safe(row['type'])]+[row[k] for k in FIELDS[1:]])
    trend=wb.create_sheet('Daily cash flow');trend.append(('Date','Receipts','Expenses'))
    for row in data['series']:
        trend.append((_safe(row['date']),row['income'],row['expenses']))
    from executive_monthly_v144 import monthly_totals
    calendar=data.get('monthly_calendar','gregorian')
    monthly=wb.create_sheet('Jalali cash flow' if calendar=='jalali' else 'Monthly cash flow')
    monthly.append(('Jalali month' if calendar=='jalali' else 'Gregorian month','Receipts','Expenses','Recorded difference'))
    for item in monthly_totals(data['series'],data.get('monthly_calendar','gregorian')):
        monthly.append((item['month'],item['income'],item['expenses'],item['difference']))
    history=data.get('multi_year',{}).get('rows',[])
    if history:
        history_sheet=wb.create_sheet('Multi-year cash flow')
        history_sheet.append(('Years ago','Start','End','Receipts','Expenses','Difference'))
        for item in history:
            history_sheet.append(tuple(item[key] for key in ('years_ago','start','end','income','expenses','difference')))
    for sheet in wb.worksheets:
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        sheet.sheet_view.showGridLines=False
        for cell in sheet[1]:
            cell.fill=PatternFill('solid',fgColor='17364B')
            cell.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
            cell.alignment=Alignment(vertical='center')
        sheet.row_dimensions[1].height=27
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.fill=PatternFill('solid',fgColor='EDF5F8' if cell.row%2==0 else 'FFFFFF')
                cell.font=Font(name='Calibri',size=10,color='18384C')
                if isinstance(cell.value,(int,float)):
                    cell.number_format='#,##0.##'
        for col in sheet.columns:
            letter=get_column_letter(col[0].column)
            longest=max((len(str(c.value or '')) for c in col),default=10)
            sheet.column_dimensions[letter].width=min(42,max(14,longest+3))
        sheet.print_options.horizontalCentered=True
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.fitToWidth=1
    temp=None
    try:
        with NamedTemporaryFile(dir=target.parent,suffix='.xlsx',delete=False) as f:temp=Path(f.name)
        wb.save(temp)
        os.replace(temp,target)
    finally:
        if temp is not None:temp.unlink(missing_ok=True)
    return target

def _pdf_font():
    """Use an installed Unicode font when available, otherwise standard Helvetica."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from pathlib import Path
    candidates=[Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
                Path('/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf'),
                Path('C:/Windows/Fonts/tahoma.ttf'),Path('C:/Windows/Fonts/arial.ttf')]
    for candidate in candidates:
        if candidate.is_file():
            try:
                if 'AshkanUnicode' not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont('AshkanUnicode',str(candidate)))
                return 'AshkanUnicode'
            except Exception:continue
    return 'Helvetica'

def _pdf_text(value,font):
    text=str(value)
    if font=='Helvetica':return text.encode('latin-1','replace').decode('latin-1')
    if any('\u0600'<=c<='\u06ff' for c in text):
        try:
            import arabic_reshaper
            from bidi.algorithm import get_display
            return get_display(arabic_reshaper.reshape(text))
        except ImportError:pass
    return text

def export_pdf(data,path):
    """Compact English-labelled PDF. Non-ASCII room names need a configured Unicode font."""
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4,landscape
    from xml.sax.saxutils import escape
    validate(data)
    target=Path(path)
    if target.suffix.lower()!='.pdf':raise ValueError('Expected .pdf output')
    temp=None
    try:
        with NamedTemporaryFile(dir=target.parent,suffix='.pdf',delete=False) as f:temp=Path(f.name)
        styles=getSampleStyleSheet(); font=_pdf_font()
        for style in styles.byName.values():style.fontName=font
        story=[Paragraph('Hotel Executive Dashboard — estimated booked room value',styles['Title']),Spacer(1,14)]
        story.append(Paragraph('Room value, ADR and RevPAR are estimates, not recognized accounting revenue.',styles['Normal']))
        table=[HEADERS]
        for row in data['room_types']:
            # Standard ReportLab fonts do not support arbitrary Persian/Unicode names.
            kind=_pdf_text(row['type'],font)
            table.append([kind]+[str(row[k]) for k in FIELDS[1:]])
        if len(table)==1:table.append(['No data']+['']*6)
        t=Table(table,repeatRows=1,colWidths=[125,50,75,110,75,80,80]);t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#17364B')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.3,colors.grey),('FONTNAME',(0,0),(-1,-1),font),('FONTSIZE',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),7)]));story.append(t)
        story.append(Spacer(1,15));story.append(Paragraph('Daily receipts and expenses',styles['Heading2']))
        daily=[('Date','Receipts','Expenses')]+[(str(r['date']),str(r['income']),str(r['expenses'])) for r in data['series']]
        dt=Table(daily,repeatRows=1);dt.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.2,colors.grey),('FONTNAME',(0,0),(-1,-1),font),('FONTNAME',(0,0),(-1,-1),font),('FONTSIZE',(0,0),(-1,-1),8)]));story.append(dt)
        from executive_monthly_v144 import monthly_totals
        story.append(Spacer(1,14))
        story.append(Paragraph('Monthly recorded cash flow (Jalali months)' if data.get('monthly_calendar')=='jalali' else 'Monthly recorded cash flow (Gregorian months)',styles['Heading2']))
        monthly_rows=[('Month','Receipts','Expenses','Difference')]
        for item in monthly_totals(data['series'],data.get('monthly_calendar','gregorian')):
            monthly_rows.append(tuple(str(item[key]) for key in ('month','income','expenses','difference')))
        monthly_table=Table(monthly_rows,repeatRows=1)
        monthly_table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#17364B')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.3,colors.grey),('FONTNAME',(0,0),(-1,-1),font),('FONTSIZE',(0,0),(-1,-1),9)]))
        story.append(monthly_table)
        history=data.get('multi_year',{}).get('rows',[])
        if history:
            story.append(Spacer(1,14))
            story.append(Paragraph('Historical recorded cash flow (same-length calendar periods)',styles['Heading2']))
            history_rows=[('Years ago','Start','End','Receipts','Expenses','Difference')]
            for item in history:
                history_rows.append(tuple(str(item[key]) for key in ('years_ago','start','end','income','expenses','difference')))
            history_table=Table(history_rows,repeatRows=1)
            history_table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#17364B')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.3,colors.grey),('FONTNAME',(0,0),(-1,-1),font),('FONTSIZE',(0,0),(-1,-1),8)]))
            story.append(history_table)
        SimpleDocTemplate(str(temp),pagesize=landscape(A4),leftMargin=20,rightMargin=20).build(story)
        os.replace(temp,target)
    finally:
        if temp is not None:temp.unlink(missing_ok=True)
    return target
