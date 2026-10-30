"""Branch-scoped executive Excel workbook. Caller supplies authorized, scoped data."""
from collections import defaultdict
from datetime import date
from pathlib import Path
from branch_csv import safe_cell


def payment_series(payments, start, end):
    """Aggregate recorded payments by calendar month within [start, end)."""
    begin, finish = date.fromisoformat(start), date.fromisoformat(end)
    if begin >= finish:
        raise ValueError('Invalid report date range')
    monthly = defaultdict(lambda: [0, 0])
    for payment in payments:
        if len(payment) != 5:
            raise ValueError('Expected five payment fields')
        created = date.fromisoformat(str(payment[4])[:10])
        if begin <= created < finish:
            key = created.strftime('%Y-%m')
            monthly[key][0] += int(payment[2])
            monthly[key][1] += 1
    return [(month, amount, count) for month, (amount, count) in sorted(monthly.items())]


def export_management_workbook(path, *, branch_name, start, end, summary, reservations, payments, language='fa'):
    """Generate bilingual Excel workbook; all supplied rows must already be access-scoped."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    if language not in ('fa', 'en'):
        raise ValueError('Unsupported language')
    if date.fromisoformat(start) >= date.fromisoformat(end):
        raise ValueError('Invalid report date range')
    reservations, payments = list(reservations), list(payments)
    if any(len(r) != 5 for r in reservations + payments):
        raise ValueError('Invalid report row')
    fa = language == 'fa'
    wb = Workbook()
    overview = wb.active
    overview.title = 'خلاصه مدیریتی' if fa else 'Executive Summary'
    bookings = wb.create_sheet('رزروها' if fa else 'Reservations')
    transactions = wb.create_sheet('پرداخت‌ها' if fa else 'Payments')
    trend = wb.create_sheet('روند درآمد' if fa else 'Revenue Trend')
    for ws in wb:
        ws.sheet_view.rightToLeft = fa
        ws.freeze_panes = 'A5' if ws != overview else 'A4'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    def title(ws, name, last_col):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
        c = ws.cell(1, 1, name)
        c.fill = PatternFill('solid', fgColor='142D4E')
        c.font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
        c.alignment = Alignment(horizontal='right' if fa else 'left', vertical='center')
        ws.row_dimensions[1].height = 34
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        ws.cell(2, 1, f'{safe_cell(branch_name)} | {start} — {end} (end exclusive)')
        ws.row_dimensions[2].height = 25
    def header(ws, names):
        for i, name in enumerate(names, 1):
            c = ws.cell(4, i, name)
            c.fill = PatternFill('solid', fgColor='087E8B')
            c.font = Font(bold=True, color='FFFFFF')
            c.alignment = Alignment(horizontal='center')
        ws.row_dimensions[4].height = 25
    def rows(ws, data, numeric_cols=()):
        for row_index, record in enumerate(data, 5):
            for col, value in enumerate(record, 1):
                c = ws.cell(row_index, col)
                if col in numeric_cols:
                    c.value = int(value)
                    c.number_format = '#,##0'
                else:
                    c.value = safe_cell(value)
                    c.data_type = 's'
                if row_index % 2 == 0:
                    c.fill = PatternFill('solid', fgColor='EAF4F5')
        ws.auto_filter.ref = f'A4:{get_column_letter(len(data[0]) if data else ws.max_column)}{max(4, 4 + len(data))}'
    title(overview, 'گزارش مدیریتی شعبه' if fa else 'Branch Management Report', 2)
    header(overview, ('شاخص', 'مقدار') if fa else ('Metric', 'Value'))
    metrics = [('اتاق‌های فعال' if fa else 'Active rooms', summary['rooms']),
               ('تعداد رزرو' if fa else 'Reservations', summary['reservations']),
               ('شب‌اتاق اشغال' if fa else 'Occupied room-nights', summary['occupied_nights']),
               ('دریافتی ثبت‌شده' if fa else 'Recorded payments', summary['payments'])]
    rows(overview, metrics, (2,))
    title(bookings, 'رزروهای شعبه' if fa else 'Branch Reservations', 5)
    header(bookings, ('شناسه رزرو','اتاق','مهمان','ورود','خروج') if fa else ('Reservation ID','Room','Guest','Check-in','Check-out'))
    rows(bookings, reservations)
    title(transactions, 'پرداخت‌های شعبه' if fa else 'Branch Payments', 5)
    header(transactions, ('شناسه پرداخت','شناسه رزرو','مبلغ','روش پرداخت','تاریخ ثبت') if fa else ('Payment ID','Reservation ID','Amount','Method','Recorded at'))
    rows(transactions, payments, (3,))
    title(trend, 'روند دریافتی ماهانه' if fa else 'Monthly Payment Trend', 3)
    header(trend, ('ماه','دریافتی','تعداد تراکنش') if fa else ('Month','Payments','Transactions'))
    monthly = payment_series(payments, start, end)
    rows(trend, monthly, (2, 3))
    if monthly:
        chart = BarChart()
        chart.title = 'دریافتی ماهانه' if fa else 'Monthly Payments'
        chart.y_axis.title = 'تومان' if fa else 'Toman'
        chart.x_axis.title = 'ماه' if fa else 'Month'
        chart.add_data(Reference(trend, min_col=2, min_row=4, max_row=4+len(monthly)), titles_from_data=True)
        chart.set_categories(Reference(trend, min_col=1, min_row=5, max_row=4+len(monthly)))
        chart.width = 18
        chart.height = 9
        trend.add_chart(chart, 'E5')
    for ws in wb:
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 26 if col != 3 else 34
    output = Path(path)
    wb.save(output)
    return output
