"""Accommodation-value estimates based on stored booking quotes, not earned revenue.

Quotes are spread evenly across the full reservation; use as an indicative KPI only.
Date intervals are half-open, and every query is branch scoped.
"""
from datetime import date
from decimal import Decimal, ROUND_HALF_UP


def allocated_quote(total, stay_start, stay_end, report_start, report_end):
    dates = [date.fromisoformat(x) for x in (stay_start, stay_end, report_start, report_end)]
    a, b, start, end = dates
    if b <= a or end <= start or total < 0:
        raise ValueError('Invalid date interval or accommodation quote')
    nights = (b-a).days
    overlap = max(0, (min(b,end)-max(a,start)).days)
    return Decimal(total) * Decimal(overlap) / Decimal(nights), overlap


def estimated_branch_performance(conn, *, branch_id, start, end, user_id, role):
    from branch_access import require_branch
    require_branch(conn, user_id, role, branch_id)
    a,b=date.fromisoformat(start),date.fromisoformat(end)
    if b<=a:raise ValueError('Invalid reporting interval')
    active_rooms=conn.execute("SELECT COUNT(*) FROM rooms WHERE branch_id=? AND status='فعال'",(branch_id,)).fetchone()[0]
    rows=conn.execute("""SELECT r.checkin,r.checkout,br.accommodation_total
      FROM reservations r LEFT JOIN booking_rates br ON br.reservation_id=r.id
      WHERE r.branch_id=? AND r.status<>'لغو شده' AND r.checkin<? AND r.checkout>?""",(branch_id,end,start)).fetchall()
    occupied=0; estimated=Decimal(0); quoted=0; missing=0
    for checkin,checkout,amount in rows:
        nights=max(0,(min(date.fromisoformat(checkout),b)-max(date.fromisoformat(checkin),a)).days)
        occupied+=nights
        if amount is None:missing+=1;continue
        portion,_=allocated_quote(amount,checkin,checkout,start,end)
        estimated+=portion;quoted+=1
    capacity=active_rooms*(b-a).days
    # Current active rooms are not a historical capacity snapshot.
    return dict(branch_id=branch_id,rooms=active_rooms,available_room_nights=capacity,
        occupied_room_nights=occupied,estimated_accommodation_value=int(estimated.quantize(Decimal('1'),rounding=ROUND_HALF_UP)),
        quoted_reservations=quoted,missing_quotes=missing,
        estimated_adr=float((estimated/occupied).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)) if occupied and not missing else None,
        estimated_revpar=float((estimated/capacity).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)) if capacity and not missing else None,
        warning='Indicative only: quotes allocated evenly; current room inventory; missing quotes excluded. Not recognized revenue.')


def export_estimated_performance(path, rows, *, start, end, language='fa'):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from branch_csv import safe_cell
    if language not in ('fa','en'):raise ValueError('Unsupported language')
    if date.fromisoformat(start)>=date.fromisoformat(end):raise ValueError('Invalid reporting interval')
    fa=language=='fa';wb=Workbook();ws=wb.active
    ws.title='شاخص‌های اقامت' if fa else 'Stay Performance';ws.sheet_view.rightToLeft=fa
    ws.append(['شاخص‌های تخمینی اقامت' if fa else 'Indicative Accommodation KPIs'])
    ws.append([f'{start} — {end}'])
    ws.append(['هشدار: ارقام تخمینی هستند؛ درآمد شناسایی‌شده نیستند. نرخ‌ها به‌طور مساوی روی شب‌های اقامت توزیع شده‌اند.' if fa else 'WARNING: Indicative only; not recognized revenue. Quotes allocated evenly across stay nights.'])
    headers=['شعبه','اتاق فعال فعلی','شب‌اتاق اشغال','شب‌اتاق ظرفیت','ارزش تخمینی اقامت','ADR تخمینی','RevPAR تخمینی','رزرو بدون نرخ'] if fa else ['Branch','Current active rooms','Occupied room-nights','Capacity room-nights','Estimated stay value','Indicative ADR','Indicative RevPAR','Missing quotes']
    ws.append(headers)
    for i,row in enumerate(rows,5):
        values=[safe_cell(row['branch_name']),row['rooms'],row['occupied_room_nights'],row['available_room_nights'],row['estimated_accommodation_value'],row['estimated_adr'],row['estimated_revpar'],row['missing_quotes']]
        for col,value in enumerate(values,1):
            cell=ws.cell(i,col,value)
            if col==1:cell.data_type='s'
            if col in (5,6,7):cell.number_format='#,##0.00'
            if i%2==0:cell.fill=PatternFill('solid',fgColor='EAF4F5')
    for cell in ws[4]:cell.fill=PatternFill('solid',fgColor='087E8B');cell.font=Font(color='FFFFFF',bold=True)
    ws['A1'].font=Font(bold=True,size=16,color='142D4E')
    ws.freeze_panes='B5';ws.auto_filter.ref=f'A4:H{max(4,4+len(rows))}'
    for col in 'ABCDEFGH':ws.column_dimensions[col].width=25
    ws.column_dimensions['A'].width=34
    wb.save(path)
    return path
