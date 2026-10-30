"""Pure aggregation and export helpers for branch-scoped transfer history."""
from collections import Counter
from datetime import date

def summarize_transfers(rows):
    """Aggregate only the rows already authorized and filtered by the caller."""
    records = list(rows)
    by_actor = Counter(str(r[5]) for r in records)
    by_source = Counter(str(r[3]) for r in records)
    by_destination = Counter(str(r[4]) for r in records)
    by_month = Counter(str(r[1])[:7] for r in records if r[1])
    return {'total': len(records), 'by_actor': dict(by_actor.most_common()),
            'by_source': dict(by_source.most_common()),
            'by_destination': dict(by_destination.most_common()),
            'by_month': dict(sorted(by_month.items()))}

def export_transfers_excel(path, rows, *, language='fa'):
    """Write a readable Excel workbook without changing the source database."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, Reference
    data = list(rows)
    summary = summarize_transfers(data)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Transfers'
    fa = language == 'fa'
    ws.sheet_view.rightToLeft = fa
    headers = ['شناسه', 'زمان', 'شماره رزرو', 'اتاق مبدأ', 'اتاق مقصد', 'کاربر'] if fa else ['ID', 'Time', 'Reservation', 'From room', 'To room', 'Actor']
    ws.append(headers)
    for record in data:
        ws.append(list(record))
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='18324D')
        cell.alignment = Alignment(horizontal='center')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:F{max(2,ws.max_row)}'
    for i, width in enumerate([13, 23, 17, 19, 19, 22], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    stats = wb.create_sheet('Summary')
    stats.sheet_view.rightToLeft = fa
    stats.append(['شاخص' if fa else 'Metric', 'مقدار' if fa else 'Value'])
    stats.append(['کل انتقال‌ها' if fa else 'Total transfers', summary['total']])
    stats.append([])
    stats.append(['ماه' if fa else 'Month', 'تعداد' if fa else 'Count'])
    for month, count in summary['by_month'].items():
        stats.append([month, count])
    stats.column_dimensions['A'].width = 28
    stats.column_dimensions['B'].width = 18
    for row_number in (1, 4):
        for cell in stats[row_number]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='18324D')
    if summary['by_month']:
        chart = BarChart()
        chart.title = 'Monthly transfers'
        chart.add_data(Reference(stats, min_col=2, min_row=4, max_row=4+len(summary['by_month'])), titles_from_data=True)
        chart.set_categories(Reference(stats, min_col=1, min_row=5, max_row=4+len(summary['by_month'])))
        stats.add_chart(chart, 'D4')
    wb.save(path)
    return summary
