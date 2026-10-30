"""Excel export for already-authorized monthly cashflow trends."""
from pathlib import Path
from datetime import date


def validate_trend_rows(rows):
    result=[]
    seen=set()
    for row in rows:
        month=row['month']
        if not isinstance(month,str) or date.fromisoformat(month+'-01').strftime('%Y-%m')!=month or month in seen:
            raise ValueError('Invalid or duplicate month')
        seen.add(month)
        values=[]
        for key in ('receipts','expenses','net'):
            value=row[key]
            if isinstance(value,bool) or not isinstance(value,int):
                raise ValueError('Financial amounts must be integers')
            values.append(value)
        if values[0]<0 or values[1]<0 or values[2]!=values[0]-values[1]:
            raise ValueError('Inconsistent cashflow')
        result.append((month,*values))
    return sorted(result)


def export_trend_excel(path,rows,language='fa',scope=''):
    if language not in ('fa','en'):raise ValueError('Unsupported language')
    records=validate_trend_rows(rows)
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.chart import LineChart,Reference
    wb=Workbook();ws=wb.active
    ws.title='روند ماهانه' if language=='fa' else 'Monthly Trend'
    ws.sheet_view.rightToLeft=language=='fa'
    ws.append(['ماه','دریافتی','هزینه','خالص'] if language=='fa' else ['Month','Receipts','Expenses','Net'])
    for month,receipts,expenses,net in records:ws.append([month,receipts,expenses,net])
    for cell in ws[1]:
        cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='17354F')
        cell.alignment=Alignment(horizontal='center')
    for col in 'ABCD':ws.column_dimensions[col].width=23
    ws.freeze_panes='A2'
    ws.auto_filter.ref=f'A1:D{max(2,len(records)+1)}'
    for row in ws.iter_rows(min_row=2):
        for cell in row[1:]:cell.number_format='#,##0;[Red](#,##0)'
    if records:
        chart=LineChart();chart.title='Monthly cash movement' if language=='en' else 'Monthly cash movement'
        chart.y_axis.title='Amount';chart.x_axis.title='Month'
        chart.add_data(Reference(ws,min_col=2,max_col=4,min_row=1,max_row=len(records)+1),titles_from_data=True)
        chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=len(records)+1))
        chart.width=19;chart.height=10
        ws.add_chart(chart,'F2')
    path=Path(path);wb.save(path)
    return path
