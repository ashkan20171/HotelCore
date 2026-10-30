"""Branch-scoped cash movement summary; NOT accrual P&L or recognized revenue."""
from datetime import date
from decimal import Decimal


def migrate(conn):
    conn.execute("""CREATE TABLE IF NOT EXISTS branch_expenses (
        id INTEGER PRIMARY KEY, branch_id INTEGER NOT NULL REFERENCES branches(id),
        category TEXT NOT NULL, description TEXT NOT NULL DEFAULT '',
        amount INTEGER NOT NULL CHECK(amount > 0),
        spent_on TEXT NOT NULL, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
        created_by INTEGER NOT NULL REFERENCES users(id))""")
    conn.execute('CREATE INDEX IF NOT EXISTS idx_branch_expenses_date ON branch_expenses(branch_id,spent_on)')


def record_expense(conn, *, branch_id, user_id, role, category, description, amount, spent_on):
    from branch_access import require_branch
    require_branch(conn,user_id,role,branch_id)
    if not category or not category.strip() or not isinstance(amount,int) or isinstance(amount,bool) or amount<=0:
        raise ValueError('Expense category and positive integer amount required')
    date.fromisoformat(spent_on)
    cursor=conn.execute('INSERT INTO branch_expenses(branch_id,category,description,amount,spent_on,created_by) VALUES(?,?,?,?,?,?)',
                        (branch_id,category.strip(),description.strip(),amount,spent_on,user_id))
    return cursor.lastrowid


def cashflow(conn, *, branch_id, user_id, role, start, end):
    from branch_access import require_branch
    require_branch(conn,user_id,role,branch_id)
    if date.fromisoformat(start)>=date.fromisoformat(end):raise ValueError('Invalid interval')
    receipts=conn.execute("""SELECT COALESCE(SUM(p.amount),0),COUNT(*) FROM payments p
        JOIN reservations r ON r.id=p.reservation_id
        WHERE r.branch_id=? AND p.created_at>=? AND p.created_at<?""",(branch_id,start,end)).fetchone()
    expenses=conn.execute("""SELECT COALESCE(SUM(amount),0),COUNT(*) FROM branch_expenses
        WHERE branch_id=? AND spent_on>=? AND spent_on<?""",(branch_id,start,end)).fetchone()
    rows=conn.execute("""SELECT spent_on,category,description,amount FROM branch_expenses
        WHERE branch_id=? AND spent_on>=? AND spent_on<? ORDER BY spent_on,id""",(branch_id,start,end)).fetchall()
    return {'branch_id':branch_id,'receipts':receipts[0],'payment_count':receipts[1],
            'expenses':expenses[0],'expense_count':expenses[1],
            'net_cash_movement':receipts[0]-expenses[0], 'expense_rows':rows}


def export_cashflow(path, report, *, start, end, branch_name, language='fa'):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from branch_csv import safe_cell
    if language not in ('fa','en'):raise ValueError('Unsupported language')
    fa=language=='fa';wb=Workbook();ws=wb.active
    ws.title='جریان نقدی' if fa else 'Cash Movement';ws.sheet_view.rightToLeft=fa
    ws.append(['گزارش گردش وجوه شعبه' if fa else 'Branch Cash Movement',safe_cell(branch_name)])
    ws.append([start,end]);ws.append(['دریافتی ثبت‌شده' if fa else 'Recorded receipts',report['receipts']])
    ws.append(['هزینه‌های تخصیص‌یافته' if fa else 'Allocated expenses',report['expenses']])
    ws.append(['خالص گردش وجوه' if fa else 'Net cash movement',report['net_cash_movement']])
    ws.append(['این گزارش سود و زیان تعهدی نیست؛ هزینه‌های قدیمی بدون شعبه شامل نمی‌شوند.' if fa else 'Not accrual profit/loss. Legacy unallocated expenses excluded.'])
    ws.append(['تاریخ' if fa else 'Date','دسته' if fa else 'Category','شرح' if fa else 'Description','مبلغ' if fa else 'Amount'])
    for day,category,description,amount in report['expense_rows']:
        ws.append([safe_cell(day),safe_cell(category),safe_cell(description),amount])
        for cell in list(ws.rows)[-1][:3]:cell.data_type='s'
    for cell in ws[7]:cell.fill=PatternFill('solid',fgColor='087E8B');cell.font=Font(color='FFFFFF',bold=True)
    ws['A1'].font=Font(size=15,bold=True,color='142D4E')
    ws.freeze_panes='A8';ws.auto_filter.ref=f'A7:D{max(7,ws.max_row)}'
    for col,width in [('A',42),('B',28),('C',46),('D',24)]:ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=8,min_col=4,max_col=4):row[0].number_format='#,##0'
    wb.save(path);return path
